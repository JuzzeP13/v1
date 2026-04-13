"""Product/services search socket processing."""

from collections import Counter
import json

from openpyxl.utils import get_column_letter

from modules.main.python.core import *
from modules.main.python.analysis import analysis_lock
from modules.common.python.config import Config
from modules.search.python.product_search_intelligence import (
    MATCH_LEVEL_ORDER,
    build_search_profile,
    build_search_queries,
    classify_product_category,
    normalize_text,
    score_result_match,
)


PRODUCTS_EXCEL_PATH = REPORTS_DIR / "products_results.xlsx"


def generate_synonyms(query):
    """Генерирует синонимы/связанные термины для предпросмотра в UI."""
    profile = build_search_profile(
        query=query or "",
        city="",
        use_ollama=False,
        manual_terms=None,
    )
    ordered = []
    seen = set()
    for item in profile.synonym_terms + profile.related_terms + [query]:
        clean = (item or "").strip()
        if not clean:
            continue
        key = clean.lower()
        if key in seen:
            continue
        seen.add(key)
        ordered.append(clean)
    return ordered[:20]


def build_product_search_queries(query: str, city: str, synonyms: list[str]):
    """Строит расширенный пул поисковых запросов с уровнями приоритета."""
    profile = build_search_profile(
        query=query,
        city=city or "",
        use_ollama=Config.PRODUCT_SEARCH_USE_OLLAMA,
        ollama_url=Config.OLLAMA_URL,
        ollama_model=Config.PRODUCT_SEARCH_OLLAMA_MODEL,
        max_ollama_variants=max(1, Config.PRODUCT_SEARCH_MAX_OLLAMA_VARIANTS),
        manual_terms=synonyms or None,
        emit=emit_status,
    )
    variants = build_search_queries(profile, max_queries=max(20, Config.PRODUCT_SEARCH_MAX_QUERIES))
    return profile, variants


def is_low_quality_product_result(url: str, title: str = "", description: str = "") -> bool:
    """Отсекает мусорные страницы для товаров/услуг."""
    haystack = f"{url} {title} {description}".lower()
    blocked = [
        "login", "signup", "register", "/cart", "/checkout", "/privacy", "/terms",
        "youtube.com", "rutube.ru", "tiktok.com", "instagram.com", "facebook.com",
        "wikipedia.org", "yandex.ru/video", "google.com/search", "captcha", "cloudflare",
    ]
    return any(item in haystack for item in blocked)


def categorize_product(title: str, description: str) -> str:
    """Определяет категорию товара/услуги по заголовку и описанию."""
    return classify_product_category(f"{title or ''} {description or ''}")


def ensure_product_results_schema() -> None:
    """Добавляет новые колонки в product_results, если БД старая."""
    try:
        with sqlite3.connect(DB_PATH) as con:
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS product_results (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    url TEXT NOT NULL,
                    title TEXT,
                    description TEXT,
                    query TEXT,
                    city TEXT,
                    category TEXT,
                    found_at TEXT,
                    analyzed_at TEXT,
                    UNIQUE(url, query, city)
                )
                """
            )
            columns = {row[1] for row in con.execute("PRAGMA table_info(product_results)").fetchall()}
            migrations = {
                "match_level": "ALTER TABLE product_results ADD COLUMN match_level TEXT DEFAULT 'partial'",
                "rank_score": "ALTER TABLE product_results ADD COLUMN rank_score REAL DEFAULT 0",
                "tags": "ALTER TABLE product_results ADD COLUMN tags TEXT",
                "related_terms": "ALTER TABLE product_results ADD COLUMN related_terms TEXT",
                "search_tokens": "ALTER TABLE product_results ADD COLUMN search_tokens TEXT",
                "source_level": "ALTER TABLE product_results ADD COLUMN source_level TEXT DEFAULT 'partial'",
            }
            for column_name, sql in migrations.items():
                if column_name not in columns:
                    con.execute(sql)
            con.execute("CREATE INDEX IF NOT EXISTS idx_product_results_rank ON product_results(rank_score DESC)")
            con.execute("CREATE INDEX IF NOT EXISTS idx_product_results_match_level ON product_results(match_level)")
            con.commit()
    except Exception as exc:
        print(f"[Product] Не удалось обновить схему product_results: {exc}")


def get_existing_product_urls() -> set:
    """Возвращает набор URL, которые уже есть в таблице product_results."""
    try:
        with sqlite3.connect(DB_PATH) as con:
            rows = con.execute("SELECT DISTINCT url FROM product_results").fetchall()
        return {r[0] for r in rows if r and r[0]}
    except Exception as e:
        print(f"[Product] Не удалось прочитать product_results: {e}")
        return set()


def _prepare_products_sheet(ws):
    headers = ["#", "Категория", "Совпадение", "Score", "URL", "Название", "Описание", "Теги", "Запрос", "Город", "Дата"]
    widths = [4, 14, 16, 8, 45, 30, 40, 26, 20, 15, 12]
    total_cols = len(headers)

    for merged in list(ws.merged_cells.ranges):
        merged_text = str(merged)
        if merged_text.startswith("A1:"):
            ws.unmerge_cells(merged_text)

    ws.merge_cells(f"A1:{get_column_letter(total_cols)}1")
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = "🛍️ TISH SEARCH — Умный поиск товаров/услуг"
    title_cell.font = Font(bold=True, size=13, color="FFFFFF")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.fill = PatternFill("solid", fgColor="CC0000")
    ws.row_dimensions[1].height = 32

    hf = PatternFill("solid", fgColor="1E1E2E")
    hfont = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, (header, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = hfont
        cell.fill = hf
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.row_dimensions[2].height = 28
    return headers


def save_products_to_excel(results: list):
    """Сохраняет результаты поиска товаров/услуг в единый Excel файл."""
    if not results:
        return

    filepath = PRODUCTS_EXCEL_PATH

    try:
        wb = openpyxl.load_workbook(str(filepath))
        if "TovaryUslugi" in wb.sheetnames:
            ws = wb["TovaryUslugi"]
        else:
            ws = wb.active
            ws.title = "TovaryUslugi"
    except FileNotFoundError:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TovaryUslugi"

    headers = _prepare_products_sheet(ws)
    total_cols = len(headers)

    fill_colors = {
        "Еда": "FFE8D6",
        "Одежда": "E8D6FF",
        "Электроника": "D6E8FF",
        "Дом и сад": "D6FFE8",
        "Авто": "FFE8E8",
        "Красота": "FFE8F0",
        "Здоровье": "E8FFE8",
        "Образование": "FFF0D6",
        "Услуги": "F0E8FF",
        "Бизнес": "FFF8D6",
        "Другое": "F0F0F0",
    }

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    ufont = Font(color="0563C1", underline="single", bold=True)

    start_row = max(3, ws.max_row + 1)

    for i, result in enumerate(results, 1):
        row = start_row + i - 1
        category = result.get("category", categorize_product(result.get("title", ""), result.get("description", "")))
        match_label = result.get("match_label") or result.get("match_level", "partial")
        rank_score = round(float(result.get("rank_score", 0) or 0), 2)
        tags_text = ", ".join(result.get("tags", [])[:8])

        ws.cell(row=row, column=1, value=row - 2).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=2, value=category).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=3, value=match_label).alignment = Alignment(horizontal="center", wrap_text=True)
        ws.cell(row=row, column=4, value=rank_score).alignment = Alignment(horizontal="center")

        url_cell = ws.cell(row=row, column=5, value=result.get("url", ""))
        url_cell.font = ufont
        url_cell.alignment = Alignment(vertical="top", wrap_text=True)

        ws.cell(row=row, column=6, value=result.get("title", "")).alignment = wrap
        ws.cell(row=row, column=7, value=result.get("description", "")).alignment = wrap
        ws.cell(row=row, column=8, value=tags_text).alignment = wrap
        ws.cell(row=row, column=9, value=result.get("query", "")).alignment = Alignment(horizontal="center", wrap_text=True)
        ws.cell(row=row, column=10, value=result.get("city", "")).alignment = Alignment(horizontal="center")
        ws.cell(
            row=row,
            column=11,
            value=result.get("found_at", "")[:10] if result.get("found_at") else datetime.now().strftime("%Y-%m-%d"),
        ).alignment = Alignment(horizontal="center")

        fill = PatternFill("solid", fgColor=fill_colors.get(category, "F0F0F0"))
        for col in range(1, total_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            cell.fill = fill

        ws.row_dimensions[row].height = 68

    wb.save(str(filepath))
    emit_status(f"📊 Excel сохранён: {filepath.name} (всего {ws.max_row - 2} записей)", "success")
    return filepath


def _serialize_list(values: list[str]) -> str:
    return json.dumps(values or [], ensure_ascii=False)


def process_product_search(query, city, max_results, synonyms):
    """Обработка умного поиска товаров/услуг."""
    if not analysis_lock.acquire(blocking=False):
        emit_status("⚠️ Анализ уже запущен!", "warn")
        state["running"] = False
        return

    ensure_product_results_schema()
    target_results = max(1, int(max_results))
    skipped_existing_db = 0
    query_errors = 0

    try:
        profile, search_queries = build_product_search_queries(query, city, synonyms)
        existing_db_urls = get_existing_product_urls()

        emit_status(
            f"🔍 Умный поиск: {len(search_queries)} вариаций | Цель: {target_results} новых URL",
            "info",
        )
        emit_status(
            f"🧠 Семантика: синонимов={len(profile.synonym_terms)}, "
            f"связанных терминов={len(profile.related_terms)}, категорий={len(profile.category_hints)}",
            "info",
        )
        if profile.ollama_queries:
            emit_status(f"🤖 Ollama-варианты: {len(profile.ollama_queries)}", "info")
        if existing_db_urls:
            emit_status(f"ℹ️ В БД уже {len(existing_db_urls)} URL товаров/услуг — они будут пропущены", "info")

        candidates_by_url: dict[str, dict] = {}
        best_by_signature: dict[str, str] = {}
        seen_urls_for_state = set()

        with DDGS() as ddgs:
            pass_multipliers = [2, 4, 6]

            for pass_idx, mult in enumerate(pass_multipliers, 1):
                if state["stop"]:
                    emit_status("⛔ Поиск остановлен пользователем", "warn")
                    break

                pass_added = 0
                emit_status(
                    f"🔄 Проход {pass_idx}/{len(pass_multipliers)} "
                    f"(x{mult}) | Кандидатов: {len(candidates_by_url)}",
                    "info",
                )

                for idx, variant in enumerate(search_queries, 1):
                    if state["stop"]:
                        break

                    # Ранний выход: если уже есть достаточно релевантных результатов
                    if len(candidates_by_url) >= target_results and pass_idx >= 2:
                        strong = sum(
                            1 for candidate in candidates_by_url.values()
                            if candidate.get("match_level") in {"exact", "synonym"}
                        )
                        if strong >= target_results:
                            emit_status("🎯 Собрано достаточно релевантных результатов, завершаю досрочно", "success")
                            break

                    search_query = variant.text
                    emit_status(
                        f"🔍 [{idx}/{len(search_queries)}] "
                        f"{variant.source_level.upper()}: «{search_query}»",
                        "info",
                    )
                    fetch_limit = max(25, min(target_results * mult, 220))

                    try:
                        results = list(ddgs.text(search_query, max_results=fetch_limit))
                    except Exception as exc:
                        query_errors += 1
                        emit_status(f"⚠️ Ошибка поиска «{search_query}»: {str(exc)[:60]}", "warn")
                        continue

                    found_for_query = 0
                    for item in results:
                        url = (item.get("href", "") or "").strip()
                        title = (item.get("title", "") or "").strip()
                        desc = (item.get("body", "") or "").strip()

                        if not url.startswith("http"):
                            continue
                        if is_junk_url(url) or is_low_quality_product_result(url, title, desc):
                            continue
                        if url in existing_db_urls:
                            skipped_existing_db += 1
                            continue

                        domain = get_domain(url)
                        if domain in {"avito.ru", "cian.ru", "hh.ru"} and url.count("/") <= 3:
                            continue

                        score_info = score_result_match(
                            title=title,
                            description=desc,
                            profile=profile,
                            query_source=variant.source_level,
                            city=city or "",
                        )
                        if score_info.get("match_level") == "none":
                            continue

                        signature = f"{domain}|{normalize_text(title)[:90]}"
                        candidate = {
                            "url": url,
                            "title": title,
                            "description": desc,
                            "query": search_query,
                            "city": city,
                            "found_at": datetime.now().isoformat(),
                            "category": score_info.get("category") or categorize_product(title, desc),
                            "match_level": score_info.get("match_level", "partial"),
                            "match_label": score_info.get("match_label", "Частичное совпадение"),
                            "rank_score": float(score_info.get("rank_score", 0.0)),
                            "tags": score_info.get("tags", []),
                            "related_terms": score_info.get("related_terms", []),
                            "matched_synonyms": score_info.get("matched_synonyms", []),
                            "source_level": variant.source_level,
                            "search_tokens": sorted(profile.query_token_set),
                        }

                        current_url_for_signature = best_by_signature.get(signature)
                        if current_url_for_signature and current_url_for_signature != url:
                            prev = candidates_by_url.get(current_url_for_signature)
                            if prev and float(prev.get("rank_score", 0.0)) >= candidate["rank_score"]:
                                continue
                            if prev:
                                candidates_by_url.pop(current_url_for_signature, None)
                        best_by_signature[signature] = url

                        existing_candidate = candidates_by_url.get(url)
                        if existing_candidate and float(existing_candidate.get("rank_score", 0.0)) >= candidate["rank_score"]:
                            continue

                        candidates_by_url[url] = candidate
                        if url not in seen_urls_for_state:
                            seen_urls_for_state.add(url)
                            state["found_urls"].append(url)

                        pass_added += 1
                        found_for_query += 1

                    if found_for_query:
                        emit_status(f"  ✅ По запросу +{found_for_query} релевантных результатов", "success")

                    time.sleep(0.2)

                if not pass_added:
                    emit_status("⚠️ На этом проходе новых релевантных ссылок не найдено", "warn")
                    break

                level_counts = Counter(
                    candidate.get("match_level", "partial") for candidate in candidates_by_url.values()
                )
                emit_status(
                    "📈 Текущее ранжирование: "
                    f"exact={level_counts.get('exact', 0)}, "
                    f"synonym={level_counts.get('synonym', 0)}, "
                    f"category={level_counts.get('category', 0)}, "
                    f"partial={level_counts.get('partial', 0)}",
                    "info",
                )

        ranked_results = sorted(
            candidates_by_url.values(),
            key=lambda item: (
                MATCH_LEVEL_ORDER.get(item.get("match_level", "none"), 99),
                -float(item.get("rank_score", 0.0)),
            ),
        )
        all_results = ranked_results[:target_results]

        if all_results:
            emit_status(
                f"🎯 Отобрано {len(all_results)} лучших результатов "
                f"из {len(candidates_by_url)} кандидатов",
                "success",
            )
            level_counts = Counter(result.get("match_level", "partial") for result in all_results)
            emit_status(
                "📊 Финальный баланс: "
                f"exact={level_counts.get('exact', 0)}, "
                f"synonym={level_counts.get('synonym', 0)}, "
                f"category={level_counts.get('category', 0)}, "
                f"partial={level_counts.get('partial', 0)}",
                "info",
            )
            print(f"[DEBUG] Сохраняю {len(all_results)} результатов...")

            # Сохраняем в БД
            now = datetime.now().isoformat()
            try:
                with sqlite3.connect(DB_PATH) as con:
                    synced = 0
                    for result in all_results:
                        con.execute(
                            """
                            INSERT INTO product_results (
                                url, title, description, query, city, category, found_at, analyzed_at,
                                match_level, rank_score, tags, related_terms, search_tokens, source_level
                            )
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                            ON CONFLICT(url, query, city) DO UPDATE SET
                                title=excluded.title,
                                description=excluded.description,
                                category=excluded.category,
                                analyzed_at=excluded.analyzed_at,
                                match_level=excluded.match_level,
                                rank_score=excluded.rank_score,
                                tags=excluded.tags,
                                related_terms=excluded.related_terms,
                                search_tokens=excluded.search_tokens,
                                source_level=excluded.source_level
                            """,
                            (
                                result.get("url", ""),
                                result.get("title", ""),
                                result.get("description", ""),
                                result.get("query", ""),
                                result.get("city", ""),
                                result.get("category", "Другое"),
                                result.get("found_at", now),
                                now,
                                result.get("match_level", "partial"),
                                float(result.get("rank_score", 0.0)),
                                _serialize_list(result.get("tags", [])),
                                _serialize_list(result.get("related_terms", [])),
                                _serialize_list(result.get("search_tokens", [])),
                                result.get("source_level", "partial"),
                            ),
                        )
                        synced += 1
                    con.commit()
                    print(f"[DEBUG] Product DB sync: synced={synced}")
            except Exception as exc:
                print(f"[ERROR] Ошибка при сохранении в БД: {exc}")
                emit_status(f"❌ Ошибка при сохранении в БД: {exc}", "error")

            # Сохраняем в Excel
            try:
                filepath = save_products_to_excel(all_results)
                emit_status(f"📊 Excel сохранён: {filepath}", "success")
            except Exception as exc:
                print(f"[ERROR] Ошибка при сохранении Excel: {exc}")
                emit_status(f"❌ Ошибка при сохранении Excel: {exc}", "error")

            # Отправляем результаты клиенту
            for idx, result in enumerate(all_results, 1):
                score_text = f"{result.get('match_label')} • score {float(result.get('rank_score', 0.0)):.1f}"
                if result.get("matched_synonyms"):
                    score_text += f" • {', '.join(result['matched_synonyms'][:3])}"
                description_preview = result.get("description", "")[:180] or "Описание отсутствует"

                socketio.emit(
                    "result",
                    {
                        "url": result.get("url", ""),
                        "type": "Товары/Услуги",
                        "category": result.get("category", "Другое"),
                        "design": f"{score_text} | {description_preview}",
                        "ux": result.get("title", "")[:200] or "Заголовок отсутствует",
                        "index": idx,
                    },
                    room=state.get("active_sid"),
                )
                state["results"].append(result)
                emit_state()

            if len(all_results) < target_results:
                emit_status(
                    f"⚠️ Собрано {len(all_results)} из {target_results} новых URL. "
                    "Часть выдачи отфильтрована как дубли/низкое качество/слабая релевантность.",
                    "warn",
                )

            emit_status(
                f"✅ Поиск завершён! Найдено {len(all_results)} новых результатов "
                f"(пропущено из БД: {skipped_existing_db}, ошибок запросов: {query_errors}).",
                "success",
            )
        else:
            emit_status("⚠️ Ничего не найдено (или всё найденное уже есть в БД)", "warn")

    except Exception as exc:
        emit_status(f"❌ Критическая ошибка поиска товаров/услуг: {str(exc)[:120]}", "error")
    finally:
        print("[DEBUG] Сброс состояния...")
        state["running"] = False
        state["phase"] = "done"
        state["city"] = ""
        state["current_url"] = ""
        emit_state()
        state["active_sid"] = None
        if analysis_lock.locked():
            try:
                analysis_lock.release()
            except RuntimeError:
                pass
